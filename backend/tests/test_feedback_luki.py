import io

import pytest
from openpyxl import load_workbook


@pytest.mark.asyncio
async def test_luki_feedback_is_listed_separately_and_exported(client, admin_headers, monkeypatch):
    """«Написать в поддержку» в «Люках САО» ведёт на /feedback?app=luki —
    администратор должен отличать такие обращения от проблем журнала обходов."""
    from app.routers import feedback
    from app.services.rate_limit import RateLimiter

    # Лимит обращений общий на весь прогон (один IP тестового клиента) —
    # свой счётчик, чтобы не отнимать попытки у остальных тестов.
    monkeypatch.setattr(feedback, "_submit_limiter", RateLimiter(max_requests=10, window_seconds=60))
    luki = await client.post("/api/v1/feedback/", json={
        "report_type": "luki", "location_text": "ОЛХ-012", "message": "Не загружается фото ПОСЛЕ",
    })
    assert luki.status_code == 201, luki.text
    assert luki.json()["report_type"] == "luki"
    app = await client.post("/api/v1/feedback/", json={
        "report_type": "app", "message": "Не открывается обход",
    })
    assert app.status_code == 201, app.text

    listed = await client.get("/api/v1/feedback/", params={"report_type": "luki"}, headers=admin_headers)
    assert listed.status_code == 200, listed.text
    ids = [item["id"] for item in listed.json()["items"]]
    assert luki.json()["id"] in ids
    assert app.json()["id"] not in ids

    exported = await client.get("/api/v1/feedback/export.xlsx", params={"report_type": "luki"}, headers=admin_headers)
    assert exported.status_code == 200, exported.text
    ws = load_workbook(io.BytesIO(exported.content))["Обращения"]
    values = {cell.value for row in ws.iter_rows(min_row=2) for cell in row}
    assert "Люки САО" in values
    assert "Не загружается фото ПОСЛЕ" in values

