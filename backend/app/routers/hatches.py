"""Журнал осмотра люков: люки площадок, их осмотр в рамках обхода,
журнал для района/округа и выгрузка в Excel.

Осмотр люка — часть обычного обхода (hatch_checks.inspection_id), права
на чтение/запись повторяют права на сам обход (get_inspection /
update_inspection в routers/inspections.py). Правила замечаний по люкам —
в app/services/hatches.py.
"""
import io
import uuid as _uuid
from datetime import date, datetime, timezone
from typing import Optional
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from geoalchemy2.elements import WKTElement
from sqlalchemy import exists, func, select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.database import get_db
from app.models import (
    Courtyard, District, Hatch, HatchCheck, Inspection, Issue, IssueStatusHistory,
    Photo, Site, User,
)
from app.schemas import (
    HatchCheckIn, HatchCheckOut, HatchCreate, HatchJournalKpis, HatchJournalOut,
    HatchJournalRow, HatchLastCheckOut, HatchOpenIssueOut, HatchOut, HatchUpdate,
    InspectionHatchOut, PhotoOut, SiteHatchOut, StatsPeriodOut,
)
from app.services.audit import log_action
from app.services.auth import get_current_user
from app.services.hatches import (
    HATCH_STATE_SHORT, HATCH_STATES, fix_state, measures_text, natural_key,
    sync_issue_for_check,
)
from app.services.permissions import check_own_or_role, in_district_scope, require_role
from app.services.statistics.definitions import OVERDUE_STATUSES
from app.services.statistics.filters import build_filter
from app.services.timezone import MSK, msk_day_bounds_utc

router = APIRouter()


# ── Вспомогательное ─────────────────────────────────────────────

def _hatch_out(h: Hatch, lat, lon) -> HatchOut:
    return HatchOut(
        id=h.id, site_id=h.site_id, number=h.number, owner=h.owner,
        location_note=h.location_note, external_id=h.external_id,
        is_active=h.is_active, created_at=h.created_at,
        lat=float(lat) if lat is not None else None,
        lon=float(lon) if lon is not None else None,
    )


def _hatch_select():
    return select(Hatch, func.ST_Y(Hatch.point).label("lat"), func.ST_X(Hatch.point).label("lon"))


def _photo_out(p: Photo) -> PhotoOut:
    return PhotoOut(
        id=p.id, target_type=p.target_type, inspection_id=p.inspection_id,
        issue_id=p.issue_id, checklist_answer_id=p.checklist_answer_id,
        url=f"/uploads/{p.storage_path}",
        thumbnail_url=f"/uploads/{p.thumbnail_path}" if p.thumbnail_path else None,
        gps_lat=p.gps_lat, gps_lon=p.gps_lon, taken_at=p.taken_at, created_at=p.created_at,
    )


def _clean(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    value = value.strip()
    return value or None


def _point(lat: Optional[float], lon: Optional[float]):
    if lat is None and lon is None:
        return None
    if lat is None or lon is None:
        raise HTTPException(422, "Координаты люка указываются парой: широта и долгота")
    return WKTElement(f"POINT({lon} {lat})", srid=4326)


async def _load_inspection(db: AsyncSession, inspection_id: UUID) -> Inspection:
    obj = (await db.execute(
        select(Inspection).where(Inspection.id == inspection_id).options(
            selectinload(Inspection.site).selectinload(Site.courtyard),
        )
    )).scalar_one_or_none()
    if not obj:
        raise HTTPException(404, "Обход не найден")
    return obj


def _inspection_district(obj: Inspection):
    return obj.site.courtyard.district_id if obj.site and obj.site.courtyard else None


def _check_view_access(user: User, obj: Inspection) -> None:
    """Как в get_inspection: чужой обход инспектор видит только в своём
    районе (не задваивать работу коллеги), проверяющий — в своей зоне."""
    if user.role == "inspector" and obj.inspector_id != user.id:
        if not in_district_scope(user, _inspection_district(obj)):
            raise HTTPException(403, "Обход вне вашего района")
    elif user.role == "reviewer":
        if not in_district_scope(user, _inspection_district(obj)):
            raise HTTPException(403, "Обход вне вашего района")


def _check_edit_access(user: User, obj: Inspection) -> None:
    """Как в update_inspection: владелец или reviewer/admin (reviewer — в
    своей зоне); проверенный обход владелец уже не меняет."""
    check_own_or_role(user, obj.inspector_id, "reviewer", "admin")
    if user.role == "reviewer" and not in_district_scope(user, _inspection_district(obj)):
        raise HTTPException(403, "Обход вне вашего района")
    is_owner = str(user.id) == str(obj.inspector_id)
    if obj.reviewed_by is not None and is_owner:
        raise HTTPException(
            409,
            "Обход уже проверен — менять осмотр люков после проверки нельзя. "
            "Дождитесь возврата на доработку от проверяющего.",
        )


async def _checks_out(db: AsyncSession, checks: list[HatchCheck]) -> dict:
    """HatchCheckOut по id люка — пакетно: имена, статусы замечаний, фото."""
    if not checks:
        return {}
    user_ids = {c.checked_by for c in checks}
    names = dict((await db.execute(
        select(User.id, User.full_name).where(User.id.in_(user_ids))
    )).all())
    issue_ids = [c.issue_id for c in checks if c.issue_id is not None]
    issues: dict = {}
    photos: dict = {}
    if issue_ids:
        for iid, status, due in (await db.execute(
            select(Issue.id, Issue.status, Issue.due_date).where(Issue.id.in_(issue_ids))
        )).all():
            issues[iid] = (status, due)
        for p in (await db.execute(
            select(Photo).where(Photo.issue_id.in_(issue_ids), Photo.target_type == "issue")
            .order_by(Photo.created_at.asc())
        )).scalars().all():
            photos.setdefault(p.issue_id, []).append(_photo_out(p))
    out = {}
    for c in checks:
        status, due = issues.get(c.issue_id, (None, None))
        out[c.hatch_id] = HatchCheckOut(
            id=c.id, inspection_id=c.inspection_id, hatch_id=c.hatch_id, state=c.state,
            fenced=c.fenced, owner_ticket=c.owner_ticket, comment=c.comment,
            issue_id=c.issue_id, issue_status=status, issue_due_date=due,
            photos=photos.get(c.issue_id, []) if c.issue_id else [],
            checked_by=c.checked_by, checked_by_name=names.get(c.checked_by),
            created_at=c.created_at, updated_at=c.updated_at,
        )
    return out


async def _inspection_hatches(db: AsyncSession, obj: Inspection) -> list[InspectionHatchOut]:
    checks = (await db.execute(
        select(HatchCheck).where(HatchCheck.inspection_id == obj.id)
    )).scalars().all()
    checked_ids = {c.hatch_id for c in checks}
    # Отключённый люк, по которому в этом обходе уже есть осмотр, остаётся
    # в списке — иначе из обхода «пропала» бы уже сделанная запись.
    visible = Hatch.is_active.is_(True)
    if checked_ids:
        visible = visible | Hatch.id.in_(checked_ids)
    rows = (await db.execute(
        _hatch_select().where(Hatch.site_id == obj.site_id, visible)
    )).all()
    by_hatch = await _checks_out(db, list(checks))
    items = [
        InspectionHatchOut(hatch=_hatch_out(r.Hatch, r.lat, r.lon), check=by_hatch.get(r.Hatch.id))
        for r in rows
    ]
    items.sort(key=lambda i: natural_key(i.hatch.number))
    return items


# ── Люки площадки (карточка площадки) ──────────────────────────

@router.get("/sites/{site_id}/hatches", response_model=list[SiteHatchOut])
async def list_site_hatches(
    site_id: UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    site = (await db.execute(
        select(Site).where(Site.id == site_id).options(selectinload(Site.courtyard))
    )).scalar_one_or_none()
    if not site:
        raise HTTPException(404, "Площадка не найдена")
    district_id = site.courtyard.district_id if site.courtyard else None
    # Как get_site: инспектор без района площадок не видит вообще.
    if current_user.role == "inspector" and current_user.district_id is None:
        raise HTTPException(403, "Нет доступа к площадкам")
    if current_user.role in ("inspector", "reviewer") and not in_district_scope(current_user, district_id):
        raise HTTPException(403, "Площадка вне вашего района")

    rows = (await db.execute(
        _hatch_select().where(Hatch.site_id == site_id, Hatch.is_active.is_(True))
    )).all()
    if not rows:
        return []
    hatch_ids = [r.Hatch.id for r in rows]

    last_checks = {}
    for check, name in (await db.execute(
        select(HatchCheck, User.full_name)
        .join(User, User.id == HatchCheck.checked_by)
        .where(HatchCheck.hatch_id.in_(hatch_ids))
        .distinct(HatchCheck.hatch_id)
        .order_by(HatchCheck.hatch_id, HatchCheck.created_at.desc())
    )).all():
        last_checks[check.hatch_id] = HatchLastCheckOut(
            state=check.state, created_at=check.created_at,
            checked_by_name=name, inspection_id=check.inspection_id,
        )

    today = datetime.now(MSK).date()
    today_start, today_end = msk_day_bounds_utc(today, today)
    checked_today = set((await db.execute(
        select(HatchCheck.hatch_id).where(
            HatchCheck.hatch_id.in_(hatch_ids),
            HatchCheck.created_at >= today_start, HatchCheck.created_at < today_end,
        )
    )).scalars().all())

    open_issues = {}
    for hatch_id, issue_id, status, due in (await db.execute(
        select(HatchCheck.hatch_id, Issue.id, Issue.status, Issue.due_date)
        .join(Issue, Issue.id == HatchCheck.issue_id)
        .where(HatchCheck.hatch_id.in_(hatch_ids), Issue.status != "closed")
        .distinct(HatchCheck.hatch_id)
        .order_by(HatchCheck.hatch_id, Issue.created_at.desc())
    )).all():
        open_issues[hatch_id] = HatchOpenIssueOut(
            id=issue_id, status=status, due_date=due,
            is_overdue=status in OVERDUE_STATUSES and due is not None and due < today,
        )

    items = [
        SiteHatchOut(
            **_hatch_out(r.Hatch, r.lat, r.lon).model_dump(),
            last_check=last_checks.get(r.Hatch.id),
            checked_today=r.Hatch.id in checked_today,
            open_issue=open_issues.get(r.Hatch.id),
        )
        for r in rows
    ]
    items.sort(key=lambda i: natural_key(i.number))
    return items


# ── Осмотр люков в обходе ──────────────────────────────────────

@router.get("/inspections/{inspection_id}/hatches", response_model=list[InspectionHatchOut])
async def list_inspection_hatches(
    inspection_id: UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    obj = await _load_inspection(db, inspection_id)
    _check_view_access(current_user, obj)
    return await _inspection_hatches(db, obj)


@router.put("/inspections/{inspection_id}/hatches/{hatch_id}", response_model=HatchCheckOut)
async def upsert_hatch_check(
    inspection_id: UUID,
    hatch_id: UUID,
    data: HatchCheckIn,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    obj = await _load_inspection(db, inspection_id)
    _check_edit_access(current_user, obj)
    hatch = await db.get(Hatch, hatch_id)
    if hatch is None:
        raise HTTPException(404, "Люк не найден")
    if hatch.site_id != obj.site_id:
        raise HTTPException(400, "Люк не относится к площадке этого обхода")
    if not hatch.is_active:
        raise HTTPException(400, "Люк отключён — отмечать его не нужно")

    # FOR UPDATE: два одновременных PUT по одному люку (двойной тап, повтор
    # после таймаута на плохой связи) не должны оба увидеть issue_id=NULL
    # и завести два замечания на один дефект.
    stmt = (
        select(HatchCheck)
        .where(HatchCheck.inspection_id == obj.id, HatchCheck.hatch_id == hatch.id)
        .with_for_update()
        .execution_options(populate_existing=True)
    )
    check = (await db.execute(stmt)).scalar_one_or_none()
    is_new = check is None
    if check is None:
        # SAVEPOINT + повторное чтение при конфликте UNIQUE(inspection_id,
        # hatch_id) — тот же приём, что у апсерта ChecklistAnswer в
        # update_inspection: конкурент мог вставить строку между SELECT и
        # INSERT, это не ошибка, а повод обновить уже вставленную.
        try:
            async with db.begin_nested():
                check = HatchCheck(
                    inspection_id=obj.id, hatch_id=hatch.id, state=data.state,
                    checked_by=current_user.id,
                )
                db.add(check)
                await db.flush()
        except IntegrityError as e:
            if getattr(getattr(e, "orig", None), "sqlstate", None) != "23505":
                raise
            check = (await db.execute(stmt)).scalar_one()
            is_new = False

    check.state = data.state
    check.checked_by = current_user.id
    check.comment = _clean(data.comment)
    if data.state == "ok":
        check.fenced = None
        check.owner_ticket = None
    else:
        check.fenced = data.fenced
        check.owner_ticket = _clean(data.owner_ticket)
    if not is_new:
        check.updated_at = datetime.now(timezone.utc)

    await sync_issue_for_check(db, check, hatch, obj, current_user.id)
    await log_action(db, str(current_user.id), "hatch_check", "hatch_check", str(check.id), {
        "inspection_id": str(obj.id), "hatch_id": str(hatch.id), "number": hatch.number,
        "state": check.state, "issue_id": str(check.issue_id) if check.issue_id else None,
    })
    await db.commit()

    fresh = (await db.execute(
        select(HatchCheck).where(HatchCheck.id == check.id).execution_options(populate_existing=True)
    )).scalar_one()
    return (await _checks_out(db, [fresh]))[fresh.hatch_id]


@router.post("/inspections/{inspection_id}/hatches/all-ok", response_model=list[InspectionHatchOut])
async def mark_remaining_hatches_ok(
    inspection_id: UUID,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """«Остальные люки исправны» — отметить исправными только те люки,
    по которым в этом обходе ещё нет отметки; уже отмеченные (в том числе
    дефекты) не трогаются. Возвращает весь список люков обхода."""
    obj = await _load_inspection(db, inspection_id)
    _check_edit_access(current_user, obj)
    hatch_ids = (await db.execute(
        select(Hatch.id).where(Hatch.site_id == obj.site_id, Hatch.is_active.is_(True))
    )).scalars().all()
    if hatch_ids:
        now = datetime.now(timezone.utc)
        # ON CONFLICT DO NOTHING — повторное нажатие или гонка с PUT по
        # тому же люку не должны ни падать, ни перезаписывать дефект на «ок».
        result = await db.execute(
            pg_insert(HatchCheck)
            .values([
                {"id": _uuid.uuid4(), "inspection_id": obj.id, "hatch_id": hid, "state": "ok",
                 "checked_by": current_user.id, "created_at": now}
                for hid in hatch_ids
            ])
            .on_conflict_do_nothing(index_elements=["inspection_id", "hatch_id"])
        )
        if result.rowcount:
            await log_action(db, str(current_user.id), "hatch_all_ok", "inspection", str(obj.id), {
                "marked": result.rowcount,
            })
        await db.commit()
    return await _inspection_hatches(db, obj)


# ── Журнал осмотра люков (район/округ) ─────────────────────────

def _scoped(stmt, district_id, section: Optional[str]):
    if district_id:
        stmt = stmt.where(Courtyard.district_id == district_id)
    if section:
        stmt = stmt.where(Courtyard.section == section)
    return stmt


def _journal_query(f, section: Optional[str], state: Optional[str]):
    closed_at = (
        select(func.max(IssueStatusHistory.created_at))
        .where(IssueStatusHistory.issue_id == HatchCheck.issue_id,
               IssueStatusHistory.new_status == "closed")
        .correlate(HatchCheck).scalar_subquery()
    )
    has_photo = (
        exists().where(Photo.issue_id == HatchCheck.issue_id, Photo.target_type == "issue")
        .correlate(HatchCheck)
    )
    stmt = (
        select(
            HatchCheck, Hatch,
            Site.type.label("site_type"),
            Courtyard.name.label("site_address"), Courtyard.section.label("section"),
            District.name.label("district_name"),
            User.full_name.label("checked_by_name"),
            Issue.status.label("issue_status"), Issue.due_date.label("due_date"),
            closed_at.label("closed_at"), has_photo.label("has_photo"),
        )
        .join(Hatch, Hatch.id == HatchCheck.hatch_id)
        .join(Site, Site.id == Hatch.site_id)
        .join(Courtyard, Courtyard.id == Site.courtyard_id)
        .join(District, District.id == Courtyard.district_id)
        .outerjoin(User, User.id == HatchCheck.checked_by)
        .outerjoin(Issue, Issue.id == HatchCheck.issue_id)
        .where(HatchCheck.created_at >= f.start_utc, HatchCheck.created_at < f.end_utc)
    )
    stmt = _scoped(stmt, f.district_id, section)
    if state == "defects":
        stmt = stmt.where(HatchCheck.state != "ok")
    elif state:
        stmt = stmt.where(HatchCheck.state == state)
    return stmt


def _validate_state(state: Optional[str]) -> Optional[str]:
    if state and state != "defects" and state not in HATCH_STATES:
        raise HTTPException(422, "Неизвестное состояние люка")
    return state or None


def _journal_row(n: int, r, today) -> HatchJournalRow:
    c: HatchCheck = r.HatchCheck
    h: Hatch = r.Hatch
    return HatchJournalRow(
        n=n, check_id=c.id, inspection_id=c.inspection_id, created_at=c.created_at,
        district_name=r.district_name, section=r.section, site_id=h.site_id,
        site_address=r.site_address, site_type=r.site_type,
        hatch_id=h.id, hatch_number=h.number, hatch_owner=h.owner, location_note=h.location_note,
        state=c.state, has_photo=bool(r.has_photo) if c.issue_id else False,
        fenced=c.fenced, owner_ticket=c.owner_ticket, comment=c.comment,
        measures=measures_text(c), checked_by_name=r.checked_by_name,
        issue_id=c.issue_id, issue_status=r.issue_status, due_date=r.due_date,
        closed_at=r.closed_at, fix_state=fix_state(r.issue_status, r.due_date, today),
    )


async def _journal_kpis(db: AsyncSession, f, section: Optional[str]) -> HatchJournalKpis:
    today = datetime.now(MSK).date()
    today_start, today_end = msk_day_bounds_utc(today, today)

    def hatch_scope(stmt):
        stmt = (
            stmt.join(Site, Site.id == Hatch.site_id)
            .join(Courtyard, Courtyard.id == Site.courtyard_id)
        )
        return _scoped(stmt, f.district_id, section)

    active = (Hatch.is_active.is_(True), Site.is_active.is_not(False))
    total_active = (await db.execute(
        hatch_scope(select(func.count(Hatch.id)).select_from(Hatch)).where(*active)
    )).scalar_one()
    checked_today = (await db.execute(
        hatch_scope(
            select(func.count(func.distinct(HatchCheck.hatch_id)))
            .select_from(HatchCheck).join(Hatch, Hatch.id == HatchCheck.hatch_id)
        ).where(*active, HatchCheck.created_at >= today_start, HatchCheck.created_at < today_end)
    )).scalar_one()
    defects = (await db.execute(
        hatch_scope(
            select(func.count(HatchCheck.id))
            .select_from(HatchCheck).join(Hatch, Hatch.id == HatchCheck.hatch_id)
        ).where(
            HatchCheck.state != "ok",
            HatchCheck.created_at >= f.start_utc, HatchCheck.created_at < f.end_utc,
        )
    )).scalar_one()
    issues_base = hatch_scope(
        select(func.count(func.distinct(Issue.id)))
        .select_from(HatchCheck)
        .join(Issue, Issue.id == HatchCheck.issue_id)
        .join(Hatch, Hatch.id == HatchCheck.hatch_id)
    ).where(Issue.status != "closed")
    not_fixed = (await db.execute(issues_base)).scalar_one()
    overdue = (await db.execute(
        issues_base.where(
            Issue.status.in_(OVERDUE_STATUSES),
            Issue.due_date.is_not(None), Issue.due_date < today,
        )
    )).scalar_one()
    return HatchJournalKpis(
        checked_today=checked_today, total_active_hatches=total_active,
        defects_in_period=defects, not_fixed=not_fixed, overdue=overdue,
    )


@router.get("/hatches/journal", response_model=HatchJournalOut)
async def hatch_journal(
    district_id: Optional[UUID] = Query(None),
    section: Optional[str] = Query(None),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    state: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("reviewer", "admin")),
):
    f = build_filter(current_user, date_from, date_to, district_id)
    state = _validate_state(state)
    section = _clean(section)
    base = _journal_query(f, section, state)
    total = (await db.execute(select(func.count()).select_from(base.subquery()))).scalar_one()
    offset = (page - 1) * page_size
    rows = (await db.execute(
        base.order_by(HatchCheck.created_at.desc(), HatchCheck.id.desc()).offset(offset).limit(page_size)
    )).all()
    today = datetime.now(MSK).date()

    sections: list[str] = []
    if f.district_id:
        sections = sorted(
            (await db.execute(
                select(Courtyard.section).where(
                    Courtyard.district_id == f.district_id, Courtyard.section.is_not(None),
                ).distinct()
            )).scalars().all(),
            key=natural_key,
        )

    return HatchJournalOut(
        period=StatsPeriodOut(date_from=f.date_from, date_to=f.date_to),
        generated_at=datetime.now(timezone.utc),
        kpis=await _journal_kpis(db, f, section),
        sections=sections,
        total=total, page=page, page_size=page_size,
        rows=[_journal_row(offset + i + 1, r, today) for i, r in enumerate(rows)],
    )


def _msk(dt: Optional[datetime], fmt: str) -> str:
    return dt.astimezone(MSK).strftime(fmt) if dt else ""


def _fix_text(row: HatchJournalRow) -> str:
    due = row.due_date.strftime("%d.%m.%Y") if row.due_date else ""
    if row.fix_state == "accepted":
        return f"Принято {_msk(row.closed_at, '%d.%m.%Y')}".strip()
    if row.fix_state == "on_check":
        return "Устранено, на проверке"
    if row.fix_state == "overdue":
        return f"Просрочено (срок {due})"
    if row.fix_state == "in_work":
        return f"В работе, срок {due}" if due else "В работе"
    return "—"


JOURNAL_XLSX_HEADERS = [
    "№", "Дата, время", "Район", "Адрес площадки", "Тип", "Люк / владелец",
    "Состояние", "Фото", "Принятые меры", "Осмотрел", "Устранено",
]


@router.get("/hatches/journal.xlsx")
async def hatch_journal_xlsx(
    district_id: Optional[UUID] = Query(None),
    section: Optional[str] = Query(None),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    state: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("reviewer", "admin")),
):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from app.services.xlsx_style import safe_append, style_data_row, style_header_row

    f = build_filter(current_user, date_from, date_to, district_id)
    state = _validate_state(state)
    section = _clean(section)
    # По возрастанию времени, как в бумажном журнале: № 1 — первая запись периода.
    rows = (await db.execute(
        _journal_query(f, section, state).order_by(HatchCheck.created_at.asc(), HatchCheck.id.asc())
    )).all()
    today = datetime.now(MSK).date()

    wb = Workbook()
    ws = wb.active
    ws.title = "Журнал осмотра люков"
    ws.append(JOURNAL_XLSX_HEADERS)
    style_header_row(ws, 1, len(JOURNAL_XLSX_HEADERS))
    for i, r in enumerate(rows, start=1):
        row = _journal_row(i, r, today)
        hatch = f"№{row.hatch_number}" + (f" · {row.hatch_owner}" if row.hatch_owner else "")
        if row.location_note:
            hatch += f" ({row.location_note})"
        address = row.site_address + (f" · {row.section}" if row.section else "")
        # safe_append — владелец люка, № заявки, комментарии и адреса вводят
        # люди (формула-инъекция, см. app/services/xlsx_style.py).
        safe_append(ws, [
            row.n,
            _msk(row.created_at, "%d.%m.%Y %H:%M"),
            row.district_name,
            address,
            row.site_type,
            hatch,
            HATCH_STATE_SHORT.get(row.state, row.state),
            "—" if row.state == "ok" else ("есть" if row.has_photo else "нет"),
            row.measures or "—",
            row.checked_by_name or "",
            _fix_text(row),
        ])
        style_data_row(ws, ws.max_row, len(JOURNAL_XLSX_HEADERS))
    for col, width in enumerate([6, 17, 20, 40, 20, 30, 22, 8, 40, 26, 26], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    filename = f"zhurnal_lyukov_{f.date_from.isoformat()}_{f.date_to.isoformat()}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ── Управление люками (админ, до получения окружного перечня) ────

@router.get("/hatches", response_model=list[HatchOut])
async def admin_list_hatches(
    site_id: UUID = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("admin")),
):
    rows = (await db.execute(_hatch_select().where(Hatch.site_id == site_id))).all()
    items = [_hatch_out(r.Hatch, r.lat, r.lon) for r in rows]
    items.sort(key=lambda h: natural_key(h.number))
    return items


async def _ensure_unique(db: AsyncSession, site_id, number: Optional[str], external_id: Optional[str], exclude_id=None):
    if number is not None:
        q = select(Hatch.id).where(Hatch.site_id == site_id, Hatch.number == number)
        if exclude_id:
            q = q.where(Hatch.id != exclude_id)
        if (await db.execute(q)).first():
            raise HTTPException(409, f"Люк №{number} на этой площадке уже есть")
    if external_id is not None:
        q = select(Hatch.id).where(Hatch.external_id == external_id)
        if exclude_id:
            q = q.where(Hatch.id != exclude_id)
        if (await db.execute(q)).first():
            raise HTTPException(409, "Люк с таким внешним идентификатором уже есть")


async def _hatch_by_id(db: AsyncSession, hatch_id) -> HatchOut:
    r = (await db.execute(_hatch_select().where(Hatch.id == hatch_id))).one()
    return _hatch_out(r.Hatch, r.lat, r.lon)


@router.post("/hatches", response_model=HatchOut)
async def admin_create_hatch(
    data: HatchCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("admin")),
):
    if await db.get(Site, data.site_id) is None:
        raise HTTPException(404, "Площадка не найдена")
    external_id = _clean(data.external_id)
    await _ensure_unique(db, data.site_id, data.number, external_id)
    hatch = Hatch(
        site_id=data.site_id, number=data.number, owner=_clean(data.owner),
        location_note=_clean(data.location_note), external_id=external_id,
        point=_point(data.lat, data.lon), is_active=True,
    )
    db.add(hatch)
    try:
        await db.flush()
    except IntegrityError:
        raise HTTPException(409, "Такой люк уже есть")
    await log_action(db, str(current_user.id), "hatch_create", "hatch", str(hatch.id), {
        "site_id": str(data.site_id), "number": data.number,
    })
    await db.commit()
    return await _hatch_by_id(db, hatch.id)


@router.patch("/hatches/{hatch_id}", response_model=HatchOut)
async def admin_update_hatch(
    hatch_id: UUID,
    data: HatchUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role("admin")),
):
    hatch = await db.get(Hatch, hatch_id)
    if hatch is None:
        raise HTTPException(404, "Люк не найден")
    fields = data.model_fields_set
    if "number" in fields and data.number is None:
        raise HTTPException(422, "Номер люка обязателен")
    new_external = _clean(data.external_id) if "external_id" in fields else None
    await _ensure_unique(
        db, hatch.site_id,
        data.number if "number" in fields else None,
        new_external,
        exclude_id=hatch.id,
    )
    if "number" in fields:
        hatch.number = data.number
    if "owner" in fields:
        hatch.owner = _clean(data.owner)
    if "location_note" in fields:
        hatch.location_note = _clean(data.location_note)
    if "external_id" in fields:
        hatch.external_id = new_external
    if "lat" in fields or "lon" in fields:
        hatch.point = _point(data.lat, data.lon)
    if "is_active" in fields and data.is_active is not None:
        hatch.is_active = data.is_active
    try:
        await db.flush()
    except IntegrityError:
        raise HTTPException(409, "Такой люк уже есть")
    await log_action(db, str(current_user.id), "hatch_update", "hatch", str(hatch.id), {
        k: getattr(data, k) for k in fields
    })
    await db.commit()
    return await _hatch_by_id(db, hatch.id)
